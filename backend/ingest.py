"""Dataset ingestion.

Two supported shapes, both required by the brief:

  1. TWO FILES   - Dataset A (previous) and Dataset B (updated) are separate
                   workbooks/sheets, each already one dataset.
  2. ONE FILE    - a single workbook holds both datasets, distinguished by a
                   "Dataset" discriminator column (this is how the reference
                   TW Impact Study workbook is built).

Reading uses python-calamine when present because it is roughly an order of
magnitude faster than openpyxl on the 70 MB reference sheets; openpyxl is the
fallback so the app still runs on a bare install.
"""

from __future__ import annotations

import os
import re
from dataclasses import dataclass, field
from typing import Any, Iterable

import numpy as np
import pandas as pd

# ----------------------------------------------------------------------------
# Backend selection
# ----------------------------------------------------------------------------

try:  # pragma: no cover - environment dependent
    import python_calamine  # noqa: F401

    _HAVE_CALAMINE = True
except Exception:  # pragma: no cover
    _HAVE_CALAMINE = False


def _engine() -> str | None:
    return "calamine" if _HAVE_CALAMINE else None


# ----------------------------------------------------------------------------
# Sheet listing
# ----------------------------------------------------------------------------


def list_sheets(path: str) -> list[str]:
    """Return the sheet names of a workbook (empty for csv/tsv)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".csv", ".tsv", ".txt"):
        return []
    if _HAVE_CALAMINE:
        try:
            from python_calamine import CalamineWorkbook

            with CalamineWorkbook.from_path(path) as wb:
                return list(wb.sheet_names)
        except Exception:
            pass
    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return list(wb.sheetnames)
        finally:
            wb.close()
    except Exception:
        return []


# ----------------------------------------------------------------------------
# Reading
# ----------------------------------------------------------------------------

_NA = {"", "na", "n/a", "null", "none", "nan", "-", "--", "#n/a", "#value!",
       "#div/0!", "#name?", "#ref!", "#num!", "#null!"}


def _is_numeric(s: pd.Series) -> bool:
    """True for genuinely numeric columns, excluding booleans.

    Deliberately uses ``is_numeric_dtype`` rather than ``dtype == object``.
    pandas 3 gives text columns dtype ``str`` (StringDtype), not ``object``, so
    an ``object`` comparison silently skips every text column - which meant the
    numeric coercion below never ran and any dataset whose numbers arrive as
    text was profiled as having no metrics at all.
    """
    return bool(pd.api.types.is_numeric_dtype(s)) and not bool(
        pd.api.types.is_bool_dtype(s))


def _clean_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Normalise headers, blank out placeholders, drop empty rows, coerce numerics."""
    # Drop unnamed/empty columns
    df = df.loc[:, [c for c in df.columns if not str(c).startswith("Unnamed:")]]

    # Header hygiene: collapse whitespace, keep original as-is otherwise
    df.columns = [re.sub(r"\s+", " ", str(c)).strip() for c in df.columns]

    # 1. Turn placeholder tokens into real NaN on every non-numeric column.
    #    This has to happen before dropping empty rows: a padding row full of
    #    empty strings is not "all NaN" and would otherwise survive as data.
    for c in df.columns:
        if _is_numeric(df[c]):
            continue
        s = df[c].astype(str).str.strip()
        df[c] = s.mask(s.str.lower().isin(_NA))

    # 2. Drop rows that are now entirely empty
    df = df.dropna(how="all")

    # 3. Numeric coercion, decided on the WHOLE column rather than a head()
    #    probe, so a stray label at the end cannot be silently dropped.
    coercion_notes: list[dict] = []
    for c in df.columns:
        if _is_numeric(df[c]):
            continue
        non_null = int(df[c].notna().sum())
        if non_null == 0:
            continue
        raw = df[c].astype(str).str.replace(",", "", regex=False)
        parsed = pd.to_numeric(raw, errors="coerce")
        ok = int(parsed.notna().sum())
        if ok / non_null >= 0.95:
            lost = non_null - ok
            if lost:
                coercion_notes.append({"column": c, "values_not_numeric": lost})
            df[c] = parsed
    if coercion_notes:
        df.attrs["coercion_notes"] = coercion_notes

    return df.reset_index(drop=True)


def load_table(
    path: str,
    sheet_name: str | None = None,
    header_row: int = 1,
    max_rows: int | None = None,
) -> pd.DataFrame:
    """Load one table into a cleaned DataFrame.

    ``header_row`` is 1-based (Excel convention) so the UI can pass what a
    human sees. Falls back across engines rather than failing outright.
    """
    ext = os.path.splitext(path)[1].lower()
    skip = max(header_row - 1, 0)

    if ext in (".csv", ".tsv", ".txt"):
        sep = "\t" if ext in (".tsv", ".txt") else ","
        df = pd.read_csv(path, sep=sep, skiprows=skip, nrows=max_rows,
                         dtype=str, keep_default_na=False, engine="python")
        return _clean_frame(df)

    errors: list[str] = []

    if _HAVE_CALAMINE:
        try:
            from python_calamine import CalamineWorkbook

            with CalamineWorkbook.from_path(path) as wb:
                data = wb.get_sheet_by_name(sheet_name) if sheet_name else wb.get_sheet_by_index(0)
                rows = data.to_python(skip_empty_area=False)
            if skip:
                rows = rows[skip:]
            if max_rows:
                rows = rows[:max_rows]
            if not rows:
                return pd.DataFrame()
            width = max(len(r) for r in rows)
            rows = [list(r) + [None] * (width - len(r)) for r in rows]
            df = pd.DataFrame(rows[1:], columns=[str(h) for h in rows[0]])
            return _clean_frame(df)
        except Exception as exc:  # pragma: no cover
            errors.append(f"calamine: {exc}")

    try:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i < skip:
                    continue
                if max_rows and len(rows) >= max_rows + 1:
                    break
                rows.append(row)
        finally:
            wb.close()
        if not rows:
            return pd.DataFrame()
        width = max(len(r) for r in rows)
        rows = [list(r) + [None] * (width - len(r)) for r in rows]
        df = pd.DataFrame(rows[1:], columns=[str(h) for h in rows[0]])
        return _clean_frame(df)
    except Exception as exc:
        errors.append(f"openpyxl: {exc}")

    raise RuntimeError("Could not read '{}'. {}".format(path, " | ".join(errors)))


# ----------------------------------------------------------------------------
# Dataset splitting (single-file mode)
# ----------------------------------------------------------------------------


@dataclass
class SplitResult:
    a: pd.DataFrame
    b: pd.DataFrame
    column: str
    value_a: Any
    value_b: Any


def split_by_discriminator(
    df: pd.DataFrame, column: str, value_a: Any, value_b: Any
) -> SplitResult:
    """Split one frame into the previous (A) and updated (B) datasets."""
    a = df[df[column] == value_a].reset_index(drop=True)
    b = df[df[column] == value_b].reset_index(drop=True)
    return SplitResult(a=a, b=b, column=column, value_a=value_a, value_b=value_b)
