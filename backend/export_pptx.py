"""Per-category PowerPoint export.

The deck is built **on top of the client's own template**, which they drop into
the repository at

    templates/impact_template.pptx

(see `templates/README.md`). The template supplies the master, the theme, the
fonts and any background graphics; the deck supplies the content. If the file is
absent the export still works - it falls back to a plain 16:9 presentation - so a
missing template is never a broken run.

What the deck contains, per category:

  * an impact headline per selected metric
  * a market table **per level** - one slide for the channels, one for the
    regions, one for anything paired without a level - matching the separate
    blocks the analysis shows
  * a Manufacturer Top-10 and a Brand Top-10 per metric
  * a "what drove the change" slide per metric
  * the client entities tracked

There is deliberately **no cover slide, no chart and no QC slide**: all three were
removed on request. Charts are to be added by hand where they are wanted, and the
validation belongs to the app, not to the deck the client receives.

Every table carries black borders and a compact layout, and a negative figure is
printed in red - the same convention the workbook uses, so the two agree.
"""

from __future__ import annotations

import glob
import os
from typing import Any, Sequence

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Emu, Inches, Pt

# --- palette -----------------------------------------------------------------
INK = RGBColor(0x1F, 0x28, 0x37)
MUTED = RGBColor(0x6B, 0x72, 0x80)
GREEN = RGBColor(0x54, 0x8C, 0x3C)
BLUE = RGBColor(0x2E, 0x6D, 0xA4)
RED = RGBColor(0xC0, 0x00, 0x00)
LIGHT_GREEN = RGBColor(0xE2, 0xEF, 0xDA)
LIGHT_BLUE = RGBColor(0xDD, 0xEB, 0xF7)
GREY = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x00, 0x00, 0x00)
HILITE = RGBColor(0xFF, 0xF7, 0xE6)

SW, SH = Inches(13.333), Inches(7.5)

# Where the client's template is looked for, relative to the repository root.
TEMPLATE_REL = os.path.join("templates", "impact_template.pptx")
_TEMPLATE_ENV = "IMPACT_PPT_TEMPLATE"

DEFAULT_DISPLAY = {"unit": "auto", "scale": 1.0, "symbol": "", "decimals": 0}

# The four sides of a table's border, in the order the schema requires them to
# appear inside <a:tcPr>. Getting this wrong produces a file PowerPoint refuses.
_BORDER_TAGS = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")

# Placeholder kinds that are page furniture rather than content. A layout whose
# only placeholders are these is a perfectly good blank canvas - the stock
# "Blank" layout has all three - so they must not count when choosing one.
_CHROME_PLACEHOLDERS = {13, 15, 16}   # SLIDE_NUMBER, FOOTER, DATE


def repo_root() -> str:
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def default_template_path() -> str:
    """The template location the user was given, overridable by environment.

    If the exact name is not there but exactly one `.pptx` is sitting in
    `templates/`, that one is used. A template saved under a slightly different
    name otherwise falls through to the plain deck silently, and the user is left
    wondering why their formatting did not apply.
    """
    env = os.environ.get(_TEMPLATE_ENV)
    if env:
        return env
    exact = os.path.join(repo_root(), TEMPLATE_REL)
    if os.path.isfile(exact):
        return exact
    folder = os.path.dirname(exact)
    found = sorted(glob.glob(os.path.join(folder, "*.pptx")))
    if len(found) == 1:
        return found[0]
    return exact


def _measured_note(blk: dict) -> str:
    """Which rows the figures were measured on - stated, not implied.

    A stacked workbook carries the Total Market row *and* the channels it covers,
    so "the category total" is ambiguous: the Total Market's own rows, or the sum
    of everything in scope. The two differ by the channels the total already
    contains (~1.48x on the reference file).
    """
    if blk.get("measured_on") == "total_market":
        name = (blk.get("baseline") or {}).get("name")
        return ("Measured on the Total Market"
                + (f" · {name}" if name else "")
                + " - the channels beneath it are not added to it again.")
    return ("Measured as the sum of every market in scope, including any market "
            "that is itself a total.")


def _fmt(v, disp: dict) -> str:
    if v is None:
        return "-"
    try:
        x = float(v) / (disp.get("scale") or 1.0)
    except Exception:
        return "-"
    dp = max(0, min(int(disp.get("decimals") or 0), 6))
    sym = disp.get("symbol") or ""
    return f"{x:,.{dp}f}{sym}"


def _pct(v) -> str:
    if v is None:
        return "-"
    try:
        return f"{float(v):+.1f}%"
    except Exception:
        return "-"


def _pp(v) -> str:
    if v is None:
        return "-"
    try:
        return f"{float(v):+.1f}pp"
    except Exception:
        return "-"


def _textbox(slide, l, t, w, h, text, size=14, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def _rect(slide, l, t, w, h, fill, line=None):
    from pptx.enum.shapes import MSO_SHAPE

    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
    sh.shadow.inherit = False
    try:
        sh.adjustments[0] = 0.06
    except Exception:
        pass
    return sh


def _kpi_card(slide, l, t, w, h, label, value, sub="", value_color=INK):
    _rect(slide, l, t, w, h, GREY)
    _textbox(slide, l + Inches(0.12), t + Inches(0.08), w - Inches(0.24),
             Inches(0.26), label, size=10, color=MUTED)
    _textbox(slide, l + Inches(0.12), t + Inches(0.32), w - Inches(0.24),
             Inches(0.42), value, size=20, bold=True, color=value_color)
    if sub:
        _textbox(slide, l + Inches(0.12), t + Inches(0.78), w - Inches(0.24),
                 Inches(0.26), sub, size=9, color=MUTED)


def _title(slide, text, sub=None):
    _textbox(slide, Inches(0.5), Inches(0.26), SW - Inches(1.0), Inches(0.5),
             text, size=22, bold=True)
    if sub:
        _textbox(slide, Inches(0.5), Inches(0.78), SW - Inches(1.0), Inches(0.32),
                 sub, size=11, color=MUTED)
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.12), SW - Inches(1.0),
                                  Emu(9525))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
    line.line.fill.background(); line.shadow.inherit = False


def _cell_border(cell, color: RGBColor = BLACK, width_pt: float = 0.75) -> None:
    """Draw a black box around a table cell.

    python-pptx exposes no border API, so the four <a:ln*> elements are written
    directly. They must be the first children of <a:tcPr>, in schema order,
    ahead of the fill - hence the explicit insert positions rather than append.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for tag in _BORDER_TAGS:
        for el in tcPr.findall(qn(tag)):
            tcPr.remove(el)
    hex_color = "%02X%02X%02X" % (color[0], color[1], color[2])
    for i, tag in enumerate(_BORDER_TAGS):
        ln = OxmlElement(tag)
        ln.set("w", str(int(Pt(width_pt))))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", hex_color)
        fill.append(clr)
        ln.append(fill)
        tcPr.insert(i, ln)


def _is_negative(text: str) -> bool:
    """A figure printed in red: a real negative, not the "no value" dash."""
    t = (text or "").strip()
    return t.startswith("-") and t not in ("-", "--")


def _style_para(p, font, bold=False, color=INK, align=None):
    """Apply the font to the paragraph's **runs**, not to its default.

    `paragraph.font` writes `a:pPr/a:defRPr` - the default a run inherits *if it
    has no colour of its own*. The colour then depends on inheritance, and there
    is nothing on the run to read back and verify: a check on the cell finds no
    colour at all, so "the negatives are red" cannot be proved from the file. The
    run's own `a:rPr` is unambiguous and is what PowerPoint actually renders.
    """
    if align is not None:
        p.alignment = align
    runs = p.runs or [p.add_run()]
    for run in runs:
        run.font.size = Pt(font)
        run.font.bold = bold
        run.font.color.rgb = color


def _table(slide, headers, rows, l, t, w, h, col_w=None, font=8.5,
           highlight_rows: set[int] = frozenset(), row_h=Inches(0.24)):
    """A compact bordered table.

    Borders are black on every cell (the client's ask), the header is shaded and
    bold, numeric cells are right-aligned, and a negative figure is printed in
    red - the same convention the workbook uses, so a reader comparing the two
    does not have to re-learn the colours. The row height is pinned so a table of
    ten rows cannot grow past the slide.
    """
    n_rows, n_cols = len(rows) + 1, len(headers)
    shape = slide.shapes.add_table(n_rows, n_cols, l, t, w, h)
    tbl = shape.table
    # Turn off the built-in banding so only our own fills show through.
    tbl.first_row = False
    tbl.horz_banding = False
    if col_w:
        total = sum(col_w)
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Emu(int(w * cw / total))
    for i in range(n_rows):
        tbl.rows[i].height = row_h
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = str(htxt)
        c.fill.solid(); c.fill.fore_color.rgb = GREY
        _style_para(c.text_frame.paragraphs[0], font, bold=True, color=INK)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_top = c.margin_bottom = Emu(0)
        c.margin_left = c.margin_right = Inches(0.03)
        _cell_border(c)
    for i, row in enumerate(rows, start=1):
        for j in range(n_cols):
            val = row[j] if j < len(row) else ""
            txt = "" if val is None else str(val)
            c = tbl.cell(i, j)
            c.text = txt
            c.fill.solid()
            c.fill.fore_color.rgb = (HILITE if (i - 1) in highlight_rows else WHITE)
            _style_para(c.text_frame.paragraphs[0], font,
                        color=RED if _is_negative(txt) else INK,
                        align=None if j == 0 else PP_ALIGN.RIGHT)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_top = c.margin_bottom = Emu(0)
            c.margin_left = c.margin_right = Inches(0.03)
            _cell_border(c)
    return tbl


# ----------------------------------------------------------------------------


def _open_presentation(template: str | None):
    """Open the client's template, or a plain deck when there is none.

    The template's own slides are removed: they are the template author's sample
    content, not this study's, and leaving them in would put a stale cover or an
    example table in front of the client. Its master, layouts, theme and slide
    size are kept, which is the part that carries the formatting.
    """
    path = template or default_template_path()
    if path and os.path.isfile(path):
        prs = Presentation(path)
        _clear_slides(prs)
        return prs, path
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    return prs, None


def _clear_slides(prs) -> None:
    rId_attr = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    sldIdLst = prs.slides._sldIdLst
    for sld in list(sldIdLst):
        rId = sld.get(rId_attr)
        try:
            prs.part.drop_rel(rId)
        except Exception:
            pass
        sldIdLst.remove(sld)


def _content_placeholders(layout) -> int:
    return sum(1 for ph in layout.placeholders
               if ph.placeholder_format.type not in _CHROME_PLACEHOLDERS)


def _blank_layout(prs):
    """The layout that puts the least on the slide by itself.

    Prefer one with no *content* placeholders. Testing for "no placeholders at
    all" never matches: even the stock "Blank" layout carries date, footer and
    slide-number placeholders, so the old test fell through to the LAST layout in
    the deck - "Vertical Title and Text" - and every generated slide arrived with
    a vertical "Click to add text" box down its side.
    """
    best, best_n = None, None
    for lay in prs.slide_layouts:
        n = _content_placeholders(lay)
        if n == 0:
            return lay
        if best_n is None or n < best_n:
            best, best_n = lay, n
    return best or prs.slide_layouts[len(prs.slide_layouts) - 1]


def _new_slide(prs, layout):
    """Add a slide and strip every placeholder the layout brought with it.

    Belt and braces alongside `_blank_layout`: whatever the template's layouts
    define, the deck contains only what this module draws. A placeholder that
    survives is an empty "Click to add text" box on a client's slide.
    """
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):
        ph._element.getparent().remove(ph._element)
    return slide


def _metrics(report: dict) -> dict[str, dict]:
    metrics = report.get("metrics") or {}
    if metrics:
        return metrics
    return {report.get("metric_key") or "": {
        "key": report.get("metric_key") or "",
        "label": report.get("metric", ""),
        "growth_applicable": report.get("growth_applicable", True),
        "display": report.get("display"),
        # Kept so the market table still leads with the Total Market (and does not
        # append a duplicate sum) on the legacy single-metric path.
        "baseline": report.get("baseline") or {},
        "total": report.get("total") or {},
        "channel_block": report.get("channel_block") or [],
        "channel_level_block": report.get("channel_level_block"),
        "region_level_block": report.get("region_level_block"),
        "market_other_block": report.get("market_other_block"),
        "manufacturer_top_n": report.get("manufacturer_top_n") or [],
        "brand_top_n": report.get("brand_top_n") or [],
        "client_brands": report.get("client_brands") or [],
        "client_manufacturers": report.get("client_manufacturers") or [],
        "contributors": report.get("contributors") or [],
    }}


def _total_of(blocks, baseline_name: str) -> dict | None:
    """The row that *is* the total, found from the block itself.

    Relying only on the declared baseline meant a run whose baseline was not
    carried through printed a "Total (sum of rows above)" row directly beneath a
    row already levelled `total` - a duplicate of the Total Market the user had
    selected, which is what they asked to be rid of.
    """
    for b in blocks or []:
        if baseline_name and b.get("name") == baseline_name:
            return b
    for b in blocks or []:
        if str(b.get("level") or "").strip().lower() == "total":
            return b
    return None


def _row_name(b: dict, total: dict | None) -> str:
    """The row's label, with the Total Market marked as such.

    The workbook decorates it; the deck did not, so the leading row of a market
    table read as one more channel that happened to be highlighted.
    """
    name = str(b.get("name") or "")[:28]
    if total is not None and b is total:
        return "Total Market · " + name
    return name


def _level_tables(blk: dict) -> list[tuple[str, str, list[dict]]]:
    """The market rows grouped by level, as the analysis presents them.

    The analysis draws a separate block per level - channels, regions, and
    anything paired without a level - and the export only ever carried the flat
    list, so a study with both a channel and a region block lost the split on the
    way out.
    """
    out: list[tuple[str, str, list[dict]]] = []
    for key, title, noun in (("channel", "Channel", "Channels"),
                             ("region", "Region", "Regions"),
                             ("other", "Market", "Markets")):
        b = blk.get("market_other_block") if key == "other" \
            else blk.get(f"{key}_level_block")
        members = (b or {}).get("members") or []
        if not members:
            continue
        rows = ([b["total"]] if b.get("total") else []) + list(members)
        out.append((title, noun, rows))
    return out


def build_category_deck(report: dict, qc: dict, out_path: str,
                        meta: dict | None = None,
                        template: str | None = None) -> str:
    meta = meta or {}
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)

    prs, used_template = _open_presentation(template)
    blank = _blank_layout(prs)

    cat = report.get("category", "")
    markets = report.get("markets") or []
    mkt_label = ", ".join(markets) if markets else "All markets"
    metrics = _metrics(report)

    for key, blk in metrics.items():
        metric = blk.get("label") or key
        disp = blk.get("display") or report.get("display") or DEFAULT_DISPLAY
        _headline_slide(prs, blank, blk, cat, metric, mkt_label, disp, meta)
        _market_slides(prs, blank, blk, cat, metric, disp)
        if blk.get("manufacturer_top_n"):
            _ranked_slide(prs, blank, blk["manufacturer_top_n"], "Manufacturer",
                          "Manufacturers", cat, metric, disp)
        if blk.get("brand_top_n"):
            _ranked_slide(prs, blank, blk["brand_top_n"], "Brand", "Brands",
                          cat, metric, disp)
        _contributors_slide(prs, blank, blk, cat, metric, disp)

    primary = next(iter(metrics.values()), {})
    pdisp = primary.get("display") or report.get("display") or DEFAULT_DISPLAY
    plabel = primary.get("label") or report.get("metric", "")
    _client_slide(prs, blank, primary, cat, plabel, pdisp)

    prs.save(out_path)
    return out_path


# ----------------------------------------------------------------------------


def _headline_slide(prs, blank, blk, cat, metric, mkt_label, disp, meta):
    """The impact headline. Kept - it is the one slide that states the finding."""
    t = blk.get("total") or {}
    growth_ok = blk.get("growth_applicable", True)
    s = _new_slide(prs, blank)
    base_note = ""
    if meta.get("baseline_market"):
        base_note = (f"  |  shares vs {meta['baseline_market']}"
                     + ("" if meta.get("baseline_verified") else " (not verified)"))
    _title(s, f"Headline impact - {metric}",
           f"{cat} | {metric} | {mkt_label}{base_note}")

    cw, gap = Inches(3.0), Inches(0.2)
    x0, y0, hcard = Inches(0.5), Inches(1.6), Inches(1.25)
    _kpi_card(s, x0, y0, cw, hcard, "BEFORE - MAT TY",
              _fmt(t.get("before_current"), disp),
              f"MAT YA {_fmt(t.get('before_prior'), disp)}")
    _kpi_card(s, x0 + (cw + gap), y0, cw, hcard, "AFTER - MAT TY",
              _fmt(t.get("after_current"), disp),
              f"MAT YA {_fmt(t.get('after_prior'), disp)}")
    ls = t.get("level_shift_pp") if growth_ok else None
    if growth_ok:
        _kpi_card(s, x0 + 2 * (cw + gap), y0, cw, hcard, "LEVEL SHIFT", _pp(ls),
                  "after growth - before growth",
                  value_color=RED if (ls is not None and ls < 0) else INK)
    else:
        _kpi_card(s, x0 + 2 * (cw + gap), y0, cw, hcard, "TOP (TY)",
                  _fmt(t.get("after_current"), disp), "highest channel by TY")
    ac = t.get("abs_change")
    _kpi_card(s, x0 + 3 * (cw + gap), y0, cw, hcard, "ABSOLUTE CHANGE",
              _fmt(ac, disp),
              ("TY - YA (no growth for this metric)" if not growth_ok
               else f"on {_fmt(t.get('before_current'), disp)} base"),
              value_color=RED if (ac is not None and ac < 0) else INK)

    y1 = Inches(3.2)
    _rect(s, Inches(0.5), y1, SW - Inches(1.0), Inches(1.35), GREY)
    if growth_ok:
        _textbox(s, Inches(0.75), y1 + Inches(0.10), Inches(6), Inches(0.28),
                 "Growth in MAT TY vs MAT YA", size=11, bold=True)
        for i, (lbl, val, col) in enumerate((
                ("Previous dataset", t.get("before_growth_pct"), 0.75),
                ("Updated dataset", t.get("after_growth_pct"), 4.8),
                ("Impact on the read", ls, 8.8))):
            _textbox(s, Inches(col), y1 + Inches(0.45), Inches(3.6), Inches(0.3),
                     lbl, size=10, color=MUTED)
            _textbox(s, Inches(col), y1 + Inches(0.75), Inches(3.6), Inches(0.42),
                     _pct(val) if i < 2 else _pp(val), size=18, bold=True,
                     color=RED if (val is not None and val < 0) else INK)
    else:
        _textbox(s, Inches(0.75), y1 + Inches(0.10), Inches(9), Inches(0.28),
                 "Distribution level - impact reported as absolute change",
                 size=11, bold=True)
        _textbox(s, Inches(0.75), y1 + Inches(0.45), Inches(4), Inches(0.3),
                 "MAT YA (previous)", size=10, color=MUTED)
        _textbox(s, Inches(0.75), y1 + Inches(0.75), Inches(4), Inches(0.42),
                 _fmt(t.get("before_prior"), disp), size=18, bold=True)
        _textbox(s, Inches(4.9), y1 + Inches(0.45), Inches(4), Inches(0.3),
                 "MAT TY (updated)", size=10, color=MUTED)
        _textbox(s, Inches(4.9), y1 + Inches(0.75), Inches(4), Inches(0.42),
                 _fmt(t.get("after_current"), disp), size=18, bold=True)
        _textbox(s, Inches(8.9), y1 + Inches(0.45), Inches(3.6), Inches(0.3),
                 "Absolute change (TY - YA)", size=10, color=MUTED)
        _textbox(s, Inches(8.9), y1 + Inches(0.75), Inches(3.6), Inches(0.42),
                 _fmt(ac, disp), size=18, bold=True,
                 color=RED if (ac is not None and ac < 0) else INK)

    _textbox(s, Inches(0.5), Inches(4.75), SW - Inches(1.0), Inches(0.3),
             _measured_note(blk), size=9, color=MUTED)


def _market_slides(prs, blank, blk, cat, metric, disp):
    """One market slide **per level**, matching the analysis's separate blocks.

    A study may pair channels and regions; the analysis draws a block for each,
    and a single exported table silently merged them. Falls back to the flat
    block when no level blocks were produced (an older report).
    """
    growth_ok = blk.get("growth_applicable", True)
    baseline_name = (blk.get("baseline") or {}).get("name") or ""
    tables = _level_tables(blk)
    if not tables:
        flat = blk.get("channel_block") or []
        if not flat:
            return
        tables = [("Channel", "Channels", flat)]

    for title, noun, blocks in tables:
        total = _total_of(blocks, baseline_name)
        s = _new_slide(prs, blank)
        if growth_ok:
            _title(s, f"Market / {title.lower()} - {metric}",
                   "How each row's growth and share of the category moved "
                   "(MAT TY before vs after)")
        else:
            _title(s, f"Market / {title.lower()} - {metric}",
                   "Distribution level: MAT YA, MAT TY and the absolute change "
                   "(TY - YA)")

        rows = []
        if growth_ok:
            for b in blocks[:12]:
                rows.append([
                    _row_name(b, total),
                    _fmt((b.get("before") or {}).get("mat_ty"), disp),
                    _fmt((b.get("after") or {}).get("mat_ty"), disp),
                    _pct((b.get("before") or {}).get("growth_pct")),
                    _pct((b.get("after") or {}).get("growth_pct")),
                    _pp((b.get("level_shift") or {}).get("mat_ty_pp")),
                    f"{((b.get('contribution') or {}).get('before_share_pct') or 0):.1f}%",
                    f"{((b.get('contribution') or {}).get('after_share_pct') or 0):.1f}%",
                ])
            headers = [noun, "BEFORE MAT TY", "AFTER MAT TY", "BEFORE growth",
                       "AFTER growth", "Level shift", "Contrib before", "Contrib after"]
            col_w = [30, 12, 12, 11, 11, 11, 11, 12]
        else:
            for b in blocks[:12]:
                rows.append([
                    _row_name(b, total),
                    _fmt((b.get("before") or {}).get("mat_ya"), disp),
                    _fmt((b.get("before") or {}).get("mat_ty"), disp),
                    _fmt((b.get("after") or {}).get("mat_ya"), disp),
                    _fmt((b.get("after") or {}).get("mat_ty"), disp),
                    _fmt(b.get("abs_change"), disp),
                ])
            headers = [noun, "BEFORE MAT YA", "BEFORE MAT TY", "AFTER MAT YA",
                       "AFTER MAT TY", "Abs change (TY - YA)"]
            col_w = [32, 13, 13, 13, 13, 16]

        # With no total anywhere in the block, the sum is the only aggregate the
        # table has, so it is appended and labelled as a sum. With one, that row
        # already leads the table and repeating it here would duplicate the Total
        # Market the user selected.
        if total is None:
            t = blk.get("total") or {}
            if growth_ok:
                rows.append(["Total (sum of rows above)",
                             _fmt(t.get("before_current"), disp),
                             _fmt(t.get("after_current"), disp),
                             _pct(t.get("before_growth_pct")),
                             _pct(t.get("after_growth_pct")),
                             _pp(t.get("level_shift_pp")), "100.0%", "100.0%"])
            else:
                rows.append(["Total (sum of rows above)",
                             _fmt(t.get("before_prior"), disp),
                             _fmt(t.get("before_current"), disp),
                             _fmt(t.get("after_prior"), disp),
                             _fmt(t.get("after_current"), disp),
                             _fmt(t.get("abs_change"), disp)])

        hi = {i for i, b in enumerate(blocks[:12])
              if total is not None and b.get("name") == total.get("name")}
        _table(s, headers, rows, Inches(0.5), Inches(1.35), SW - Inches(1.0),
               Inches(0.24 * (len(rows) + 1)), col_w=col_w, font=9,
               highlight_rows=hi)
        _textbox(s, Inches(0.5), Inches(6.85), SW - Inches(1.0), Inches(0.3),
                 ("Share is of the category total in each dataset."
                  if growth_ok else
                  "Distribution values are percentages out of 100 and never exceed "
                  "100; no contribution is computed for this metric."),
                 size=9, color=MUTED)


def _ranked_slide(prs, blank, block, noun, plural, cat, metric, disp):
    s = _new_slide(prs, blank)
    _title(s, f"Top {len(block)} {plural.lower()}",
           f"The largest in the previous database, followed into the updated one "
           f"({metric})")
    rows = []
    for b in block:
        rows.append([
            b["name"][:26],
            _fmt(b.get("before_prior"), disp),
            _fmt(b.get("before_current"), disp),
            _fmt(b.get("after_prior"), disp),
            _fmt(b.get("after_current"), disp),
            _pp(b.get("share_change_pp")),
            b["rank_before"] or "-",
            b["rank_after"] or "-",
            (f"{b['rank_change']:+d}" if b.get("rank_change") is not None else "-"),
            b.get("movement") or "",
        ])
    headers = [noun, "BEFORE MAT YA", "BEFORE MAT TY", "AFTER MAT YA", "AFTER MAT TY",
               "Share chg", "Rank BEFORE", "Rank AFTER", "Rank chg", "Movement"]
    hi = {i for i, b in enumerate(block)
          if b.get("movement") in ("NEW", "EXITED")}
    _table(s, headers, rows, Inches(0.5), Inches(1.35), SW - Inches(1.0),
           Inches(0.24 * (len(rows) + 1)),
           col_w=[20, 10, 10, 10, 10, 9, 10, 10, 8, 11], font=8.5,
           highlight_rows=hi)
    _textbox(s, Inches(0.5), Inches(6.85), SW - Inches(1.0), Inches(0.3),
             "MAT YA / MAT TY are each side's own MAT. No growth column: for a ranked "
             "entity the share change is the comparable movement.",
             size=9, color=MUTED)


def _client_slide(prs, blank, blk, cat, metric, disp):
    clients = (blk.get("client_brands") or []) + (blk.get("client_manufacturers") or [])
    if not clients:
        return
    s = _new_slide(prs, blank)
    _title(s, "Client entities tracked", "Reported independently of the Top-N cut")
    rows = []
    for b in clients:
        if not b.get("found", True):
            rows.append([b.get("name", ""), "-", "-", "-", "-", "-", "not found",
                         b.get("note", "")])
            continue
        rows.append([
            b["name"][:26],
            _fmt(b.get("before_current"), disp),
            _fmt(b.get("after_current"), disp),
            b.get("rank_before") or "-", b.get("rank_after") or "-",
            _pp(b.get("share_change_pp")),
            "Yes" if b.get("in_top_n") else "No",
            b.get("movement") or "",
        ])
    _table(s, ["Client entity", "BEFORE MAT TY", "AFTER MAT TY", "Rank BEFORE",
               "Rank AFTER", "Share chg", "In Top-N", "Movement"],
           rows, Inches(0.5), Inches(1.35), SW - Inches(1.0),
           Inches(0.24 * (len(rows) + 1)),
           col_w=[26, 12, 12, 10, 10, 10, 9, 11], font=9)


def _contributors_slide(prs, blank, blk, cat, metric, disp):
    """The gainers and losers that drove the change - one slide per metric.

    Gain and Loss are the two halves of one comparison, so both are drawn the same
    way and carry the same two figures. The block exists per metric in the report,
    so it is shown per metric here too: a run that measures Value, Volume and ND
    otherwise reported "what drove the change" for Sales Value alone.
    """
    contribs = blk.get("contributors") or []
    if not contribs:
        return
    s = _new_slide(prs, blank)
    _title(s, f"What drove the change - {metric}",
           "Largest absolute gainers and losers in the category")

    y = Inches(1.35)
    for c in contribs[:2]:
        gainers = c.get("gainers") or []
        losers = c.get("losers") or []
        rows = []
        for it in gainers:
            rows.append([it["name"][:26], "Gain", _fmt(it.get("abs_change"), disp),
                         _pct(it.get("contribution_to_change_pct"))])
        for it in losers:
            rows.append([it["name"][:26], "Loss", _fmt(it.get("abs_change"), disp),
                         _pct(it.get("contribution_to_change_pct"))])
        if not rows:
            rows = [["-", "-", "-", "no measurable movement"]]
        _textbox(s, Inches(0.5), y, Inches(6), Inches(0.28),
                 f"{c.get('level', '').title()}s", size=12, bold=True)
        y += Inches(0.32)
        _table(s, ["Entity", "Direction", "Absolute change",
                   "Contribution to change"],
               rows, Inches(0.5), y, SW - Inches(1.0),
               Inches(0.24 * (len(rows) + 1)),
               col_w=[42, 14, 20, 24], font=8.5)
        y += Inches(0.24 * (len(rows) + 1) + 0.35)
        if y > Inches(6.0):
            break

    _textbox(s, Inches(0.5), Inches(6.95), SW - Inches(1.0), Inches(0.3),
             "An entity present in only one dataset has no measurable change and is "
             "reported in the Top-N movement as NEW / EXITED instead.",
             size=9, color=MUTED)
